"""Phase 3: macro ingestion (SPEC.md §6.4).

Reads the two hand-downloaded files in data/raw/macro/ (local files only, no network): BLS
CPI-U food away from home (CUSR0000SEFV) and Census food services and drinking places sales
(NAICS 722, seasonally adjusted). Checks that each file's header names the expected series,
parses whichever layout the file has (one row per year with a column per month, or one row
per month), and writes one row per (series_id, month) from 2024-01 onward to
data/interim/macro_monthly.parquet. Values are stored as published; nothing is filled,
interpolated or rounded. Then runs the Phase 3 checks M1-M6.

Run: python src/ingest_macro.py
"""

import sys
import warnings
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = ROOT / "data" / "raw" / "macro"
OUT_PATH = ROOT / "data" / "interim" / "macro_monthly.parquet"

START = pd.Timestamp("2024-01-01")
MIN_LAST = pd.Timestamp("2025-12-01")  # M3: each series must reach at least 2025-12

# Download dates, recorded by hand when each file was saved (SPEC.md §6.4: the download date is
# part of the lineage). The BLS file carries no download date of its own; the Census file's
# header line "Data Extracted on: October 24, 2026" agrees with its date.
RETRIEVED_ON = {"CUSR0000SEFV": date(2026, 10, 23), "CENSUS_722_SA": date(2026, 10, 24)}

# (series_id, file, cell texts that mean "no value published"). A blank BLS month cell has no
# value; Census prints NA, (S) or Z instead of a number.
SOURCES = [
    ("CUSR0000SEFV", "bls_CUSR0000SEFV.xlsx", {""}),
    ("CENSUS_722_SA", "census_naics722_sa.xlsx", {"NA", "(S)", "Z"}),
]

# Header lines each file must contain, exactly as the agency prints them.
EXPECTED_HEADER = {
    "CUSR0000SEFV": [["Series Id:", "CUSR0000SEFV"], ["Seasonally Adjusted"],
                     ["Area:", "U.S. city average"], ["Item:", "Food away from home"]],
    "CENSUS_722_SA": [["722: Food Services and Drinking Places: U.S. Total"],
                      ["Seasonally Adjusted Sales - Monthly [Millions of Dollars]"]],
}

# Months the agency did not publish. M2 accepts exactly these and nothing else; they are never
# filled. CPI 2025-10: cell K22 of the BLS file is blank, with the cell comment
# "*  Data unavailable due to the 2025 lapse in appropriations".
DOCUMENTED_GAPS = {"CUSR0000SEFV": {pd.Timestamp("2025-10-01")}}

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
IGNORED_COLUMNS = {"Annual", "HALF1", "HALF2"}  # BLS annual and half-year averages, not months
FLAG_WORDS = ["advance", "preliminary", "revised"]

COLUMNS = ["series_id", "date", "value", "retrieved_on"]

warnings.filterwarnings("ignore", message="Workbook contains no default style")


# ---------- parse ----------

def text(v):
    return "" if v is None else str(v).strip()


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def read_sheet(path):
    """The workbook's only sheet. Not opened read_only: the BLS file's stored sheet dimension
    (A1:A1) is wrong, and read_only mode trusts it."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if len(wb.worksheets) != 1:
        sys.exit(f"{path.name}: expected 1 sheet, found {wb.sheetnames}. Stopping.")
    return wb.worksheets[0]


def find_table(ws, path):
    """(layout, header row number). Exactly one table header is allowed per file."""
    found = []
    for row in ws.iter_rows():
        cells = [text(c.value) for c in row]
        if cells[:13] == ["Year"] + MONTHS:
            found.append(("one row per year, one column per month", row[0].row))
        elif cells[:2] == ["Period", "Value"]:
            found.append(("one row per month", row[0].row))
    if len(found) != 1:
        sys.exit(f"{path.name}: expected exactly one table header, found {found}. Stopping.")
    return found[0]


def check_header(ws, header_row, series_id, path):
    lines = [[text(c.value) for c in r if text(c.value)] for r in ws.iter_rows(max_row=header_row - 1)]
    lines = [ln for ln in lines if ln]
    for ln in lines:
        print(f"  header: {' '.join(ln)}")
    missing = [ln for ln in EXPECTED_HEADER[series_id] if ln not in lines]
    if missing:
        sys.exit(f"{path.name}: header lines not found: {missing}. Stopping.")
    if any("not seasonally adjusted" in " ".join(ln).lower() for ln in lines):
        sys.exit(f"{path.name}: header mentions not seasonally adjusted data. Stopping.")


def table_rows(ws, header_row, path):
    """Data rows below the header, up to the first empty row. Anything after that stops."""
    rows, ended = [], False
    for r in ws.iter_rows(min_row=header_row + 1):
        empty = all(text(c.value) == "" for c in r)
        if ended and not empty:
            sys.exit(f"{path.name}: unexpected content after the table at row {r[0].row}. Stopping.")
        if empty:
            ended = True
        else:
            rows.append(r)
    return rows


def cell_value(c, missing, path):
    """(value, None) for a number, (None, raw text) for a published no-value marker; else stop."""
    if is_number(c.value):
        return float(c.value), None
    if text(c.value) in missing:
        return None, text(c.value)
    sys.exit(f"{path.name}: cell {c.coordinate} holds {c.value!r}, not a number "
             f"or a no-value marker {sorted(missing)}. Stopping.")


def parse_wide(ws, header_row, rows, missing, path):
    header = list(ws.iter_rows(min_row=header_row, max_row=header_row))[0]
    extra = [text(c.value) for c in header[13:] if text(c.value)]
    unknown = [e for e in extra if e not in IGNORED_COLUMNS]
    if unknown:
        sys.exit(f"{path.name}: unexpected columns {unknown}. Stopping.")
    print(f"  ignored columns: {extra or 'none'}")
    out = []
    for r in rows:
        y = r[0].value
        if not is_number(y) or y != int(y):
            sys.exit(f"{path.name}: cell {r[0].coordinate} holds {y!r}, not a year. Stopping.")
        for m in range(12):
            c = r[1 + m]
            v, raw = cell_value(c, missing, path)
            out.append((pd.Timestamp(int(y), m + 1, 1), v, raw, c))
    return out


def parse_long(rows, missing, path):
    out = []
    for r in rows:
        if any(text(c.value) for c in r[2:]):
            sys.exit(f"{path.name}: row {r[0].row} has cells beyond Period and Value. Stopping.")
        try:
            d = datetime.strptime(text(r[0].value), "%b-%Y")
        except ValueError:
            sys.exit(f"{path.name}: cell {r[0].coordinate} holds {r[0].value!r}, not a Mon-YYYY period. Stopping.")
        v, raw = cell_value(r[1], missing, path)
        out.append((pd.Timestamp(d), v, raw, r[1]))
    return out


def flag_cells(ws):
    """Text cells that mention advance, preliminary or revised estimates."""
    return [(c.coordinate, c.value) for r in ws.iter_rows() for c in r
            if isinstance(c.value, str) and any(w in c.value.lower() for w in FLAG_WORDS)]


def load_series(series_id, filename, missing):
    path = RAW_DIR / filename
    print(f"\n--- {series_id}: {path.relative_to(ROOT)} ---")
    ws = read_sheet(path)
    layout, header_row = find_table(ws, path)
    check_header(ws, header_row, series_id, path)
    print(f"  layout: {layout} (table header at row {header_row})")
    rows = table_rows(ws, header_row, path)
    parsed = (parse_wide(ws, header_row, rows, missing, path) if layout.startswith("one row per year")
              else parse_long(rows, missing, path))

    dropped = [(d, raw, c) for d, v, raw, c in parsed if v is None]
    print(f"  no-value cells dropped: {len(dropped)}")
    for d, raw, c in dropped:
        note = f"  comment: {c.comment.text.strip()!r}" if c.comment else ""
        print(f"    {d:%Y-%m}  cell {c.coordinate} = {c.value!r}{note}")

    kept = pd.DataFrame([(d, v) for d, v, raw, c in parsed if v is not None], columns=["date", "value"])
    before = int((kept.date < START).sum())
    kept = kept[kept.date >= START]
    print(f"  months with a value: {len(kept) + before}; kept from {START:%Y-%m} onward: {len(kept)}; "
          f"before {START:%Y-%m}, not kept: {before}")
    kept.insert(0, "series_id", series_id)
    kept["retrieved_on"] = RETRIEVED_ON[series_id]
    return kept, flag_cells(ws)


def main():
    frames, flags = [], {}
    for series_id, filename, missing in SOURCES:
        df, flags[series_id] = load_series(series_id, filename, missing)
        frames.append(df)

    out = pd.concat(frames, ignore_index=True)[COLUMNS].sort_values(["series_id", "date"]).reset_index(drop=True)
    out["date"] = out["date"].dt.date
    out["value"] = out["value"].astype("float64")
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    out.to_parquet(OUT_PATH, index=False)
    print(f"\nwrote {len(out)} rows to {OUT_PATH}")

    checks(flags)


# ---------- checks ----------

def checks(flags):
    df = pd.read_parquet(OUT_PATH)
    df["date"] = pd.to_datetime(df["date"])
    series = [s for s, _, _ in SOURCES]
    results = {}

    def verdict(name, ok, rule):
        results[name] = ok
        print(f"[{'PASS' if ok else 'FAIL'}] {name}  ({rule})")

    def months(dates):
        return ", ".join(f"{d:%Y-%m}" for d in dates) or "none"

    print("\n=== M1. Rows, first and last month, first 3 and last 3 rows per series ===")
    for s in series:
        sub = df[df.series_id == s]
        print(f"\n{s}: {len(sub)} rows, first {sub.date.min():%Y-%m}, last {sub.date.max():%Y-%m}")
        show = sub.assign(date=sub.date.dt.date)
        print(show.head(3).to_string(index=False))
        print(show.tail(3).to_string(index=False, header=False))
    first_ok = all((df.series_id == s).any() and df[df.series_id == s].date.min() == START for s in series)
    verdict("M1", first_ok, f"both series present, each starting at {START:%Y-%m}")

    print("\n=== M2. Missing months between 2024-01 and each series' last month (must be none "
          "besides documented gaps) ===")
    undocumented = {}
    for s in series:
        sub = df[df.series_id == s]
        expected = pd.date_range(START, sub.date.max(), freq="MS")
        gaps = expected.difference(pd.DatetimeIndex(sub.date))
        documented = [d for d in gaps if d in DOCUMENTED_GAPS.get(s, set())]
        undocumented[s] = [d for d in gaps if d not in DOCUMENTED_GAPS.get(s, set())]
        print(f"{s}: documented gap: {months(documented)}; other missing: {months(undocumented[s])}")
    verdict("M2", all(not u for u in undocumented.values()),
            "the only missing months are the documented gaps")

    print(f"\n=== M3. Last month with a value (must be {MIN_LAST:%Y-%m} or later) ===")
    last = {s: df[df.series_id == s].date.max() for s in series}
    for s in series:
        print(f"{s}: {last[s]:%Y-%m}")
    verdict("M3", all(d >= MIN_LAST for d in last.values()), f"every series reaches {MIN_LAST:%Y-%m}")

    print("\n=== M4. Duplicate (series_id, date) keys (must be none) ===")
    dup = df[df.duplicated(["series_id", "date"], keep=False)]
    print(dup.to_string(index=False) if len(dup) else "NONE")
    verdict("M4", dup.empty, "no duplicate (series_id, date)")

    print("\n=== M5. Values for the hand check against the agency pages (read from the parquet) ===")
    wanted = [("CUSR0000SEFV", m) for m in ["2024-06", "2024-12", "2025-06", "2025-12"]] + \
             [("CENSUS_722_SA", m) for m in ["2024-12", "2025-01", "2025-02", "2025-07", "2025-08", "2025-09"]]
    absent = []
    for s, m in wanted:
        hit = df[(df.series_id == s) & (df.date == pd.Timestamp(m + "-01"))]
        if hit.empty:
            absent.append((s, m))
        print(f"{s:<14} {m}  {float(hit.value.iloc[0])!r}" if len(hit) else f"{s:<14} {m}  NOT IN PARQUET")
    verdict("M5", not absent, "every requested month is in the parquet; the values themselves are for your hand check")

    print("\n=== M6. Census months the file marks as advance or preliminary ===")
    f = flags["CENSUS_722_SA"]
    print("\n".join(f"{coord}: {v!r}" for coord, v in f) if f else "NOT IN FILE")
    verdict("M6", True, "informational; reports what the file marks, nothing to fail")

    print(f"\n{sum(results.values())}/{len(results)} checks passed")
    if not all(results[k] for k in ("M2", "M3", "M4")):
        sys.exit("M2, M3 or M4 failed. Stopping.")
    if not all(results.values()):
        sys.exit(1)


if __name__ == "__main__":
    main()
